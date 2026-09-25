"""Gera o relatório final e seus gráficos a partir dos CSVs versionados (#40)."""
from pathlib import Path
import csv
import hashlib
import json
import sys
from statistics import median, quantiles

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'lab02/scripts'))
from analyze_rq1_rq2 import read_trials, analyze, paired_test
from static_metrics import compute_metrics

OUT = ROOT / 'relatorio/lab02'
FIG = OUT / 'figuras'
BLUE, ORANGE = '#0072B2', '#D55E00'
COLORS = {'sem_ia': BLUE, 'com_ia': ORANGE}
PARTS = [('Druitti', 'Ferreira', 'ferreira'), ('gabsant07', 'Santiago', 'santiago'),
         ('tavares', 'Tavares', 'tavares')]


def fmt(n, digits=2):
    return f'{n:.{digits}f}'.replace('.', ',')


def describe(values):
    q1, _, q3 = quantiles(values, n=4, method='inclusive')
    return {'mediana': median(values), 'q1': q1, 'q3': q3, 'iqr': q3 - q1}


def load_data():
    trials = read_trials(ROOT / 'lab02/data/trials.csv')
    with (ROOT / 'lab02/data/static_metrics.csv').open() as stream:
        static = list(csv.DictReader(stream))
    keys = lambda r: (r['integrante'], r['kata'], r['tratamento'])
    index = {keys(r): r for r in static}
    if len(index) != len(static) or set(index) != {keys(r) for r in trials}:
        raise ValueError('Métricas estáticas devem corresponder aos 18 trials, sem duplicatas.')
    folders = {p: folder for p, _, folder in PARTS}
    for row in trials:
        original = index[keys(row)]
        path = ROOT / 'lab02/trials' / folders[row['integrante']] / row['kata'] / 'kata.py'
        measured = compute_metrics(path, False)
        for key in ('loc', 'sloc', 'complexidade_media'):
            if float(original[key]) != measured[key]:
                raise ValueError(f'Métrica {key} diverge da solução preservada: {path}')
            row[key] = measured[key]
    return trials


def save_figure(fig, name):
    fig.savefig(FIG / f'{name}.png', dpi=300, facecolor='white')
    fig.savefig(FIG / f'{name}.svg', facecolor='white')
    plt.close(fig)


def clean_axis(ax):
    ax.spines[['top', 'right']].set_visible(False)
    ax.spines[['left', 'bottom']].set_color('#B0B0B0')
    ax.tick_params(length=0, pad=7)
    ax.set_axisbelow(True)


def figures(rows):
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 11,
                         'axes.labelsize': 11, 'axes.titlesize': 12,
                         'svg.fonttype': 'none'})
    fig, ax = plt.subplots(figsize=(7, 2.65))
    for y, (person, name, _) in enumerate(PARTS):
        values = {t: median(r['tempo_segundos'] for r in rows
                            if r['integrante'] == person and r['tratamento'] == t)
                  for t in COLORS}
        ax.plot([values['com_ia'], values['sem_ia']], [y, y], color='#A3A3A3', lw=1.6)
        for t, value in values.items():
            ax.scatter(value, y, color=COLORS[t], marker='o' if t == 'com_ia' else 's', s=55, zorder=3)
            ax.annotate(fmt(value) + ' s', (value, y), xytext=(4, 10),
                        textcoords='offset points', fontsize=10, color=COLORS[t])
    ax.set_yticks(range(3), [n for _, n, _ in PARTS])
    ax.set_ylim(2.4, -0.55)
    ax.set_xlim(0, 790)
    ax.set_xticks([0, 200, 400, 600])
    ax.set_xlabel('Tempo mediano por participante (segundos)')
    ax.grid(axis='x', color='#E5E5E5')
    clean_axis(ax)
    fig.legend(handles=[Line2D([], [], marker=m, linestyle='', color=c, label=l)
                        for m, c, l in [('o', ORANGE, 'Com IA'), ('s', BLUE, 'Sem IA')]],
               loc='upper center', ncol=2, frameon=False, bbox_to_anchor=(0.57, 1.02))
    fig.subplots_adjust(left=.16, right=.99, top=.78, bottom=.25)
    save_figure(fig, 'rq1_tempo_pareado')

    fig, ax = plt.subplots(figsize=(7, 1.3))
    for y, t in enumerate(('com_ia', 'sem_ia')):
        values = [r['taxa_sucesso'] * 100 for r in rows if r['tratamento'] == t]
        ax.scatter(values, [y] * len(values), color=COLORS[t],
                   marker='o' if t == 'com_ia' else 's', s=55, zorder=3)
        ax.text(97, y, '9 de 9 trials com 100%', ha='right', va='center', fontsize=10)
    ax.set_yticks([0, 1], ['Com IA', 'Sem IA'])
    ax.set_ylim(1.5, -.5)
    ax.set_xlim(0, 103)
    ax.set_xticks([0, 25, 50, 75, 100], ['0%', '25%', '50%', '75%', '100%'])
    ax.grid(axis='x', color='#E5E5E5')
    clean_axis(ax)
    fig.subplots_adjust(left=.16, right=.98, top=.95, bottom=.28)
    save_figure(fig, 'rq2_sucesso')

    fig, axes = plt.subplots(1, 2, figsize=(7, 3.3))
    for ax, metric, title, label in zip(axes, ['complexidade_media', 'loc'],
                                       ['Complexidade ciclomática', 'Tamanho do código'],
                                       ['CC média por função', 'Linhas físicas (LOC)']):
        for x, t in enumerate(('com_ia', 'sem_ia')):
            values = sorted(r[metric] for r in rows if r['tratamento'] == t)
            q = describe(values)
            ax.plot([x, x], [q['q1'], q['q3']], color=COLORS[t], lw=6, alpha=.35)
            # Deslocamento horizontal apenas para revelar observações coincidentes.
            for val in sorted(set(values)):
                count = values.count(val)
                offsets = np.linspace(-.19, .19, count) if count > 1 else [0]
                ax.scatter(x + np.array(offsets), [val] * count, color=COLORS[t],
                           marker='o' if t == 'com_ia' else 's', s=28, zorder=3)
            ax.plot([x - .24, x + .24], [q['mediana']] * 2, color='black', lw=1.8, zorder=4)
            ax.text(x, 1.02, f"Mediana {fmt(q['mediana'], 0)}", transform=ax.get_xaxis_transform(),
                    ha='center', fontsize=10)
        ax.set_title(title, pad=30, loc='left')
        ax.set_xticks([0, 1], ['Com IA', 'Sem IA'])
        ax.set_xlim(-.6, 1.6)
        ax.set_ylabel(label)
        ax.set_ylim((0, 6) if metric == 'complexidade_media' else (0, 36))
        ax.grid(axis='y', color='#E5E5E5')
        clean_axis(ax)
    fig.subplots_adjust(left=.09, right=.98, top=.75, bottom=.17, wspace=.4)
    save_figure(fig, 'rq3_estrutura')


def text(doc, content, style=None):
    return doc.add_paragraph(content, style)


def table(doc, headers, values):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.autofit = False
    for cell, name in zip(tab.rows[0].cells, headers):
        cell.text = name
    for row in values:
        for cell, value in zip(tab.add_row().cells, row):
            cell.text = str(value)
    for i, row in enumerate(tab.rows):
        for cell in row.cells:
            props = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'E8EEF2' if i == 0 else 'FFFFFF')
            props.append(shading)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                e = OxmlElement(f'w:{side}')
                e.set(qn('w:w'), '90')
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            props.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = i == 0
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}')
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + k), v)
        borders.append(e)
    tab._tbl.tblPr.append(borders)
    text(doc, '').paragraph_format.space_after = Pt(0)


def picture(doc, name, caption):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    inline = p.add_run().add_picture(str(FIG / f'{name}.png'), width=Inches(7))
    inline._inline.docPr.set('descr', caption)
    text(doc, caption, 'Caption')


def link(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run(label + ': ')
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), doc.part.relate_to(url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = url
    run.append(t)
    h.append(run)
    p._p.append(h)
    p.style = 'Caption'


def build(rows):
    summary, pairs, tests = analyze(rows)
    stats = {metric: {t: describe([r[metric] for r in rows if r['tratamento'] == t])
                      for t in COLORS} for metric in ('tempo_segundos', 'taxa_sucesso', 'complexidade_media', 'loc')}
    cc_diff = [median(r['complexidade_media'] for r in rows if r['integrante'] == p and r['tratamento'] == 'com_ia') -
               median(r['complexidade_media'] for r in rows if r['integrante'] == p and r['tratamento'] == 'sem_ia')
               for p, _, _ in PARTS]
    tests['complexidade_media'] = paired_test(cc_diff, 'two-sided')
    outliers = []
    for t in COLORS:
        s = stats['tempo_segundos'][t]
        for r in rows:
            if r['tratamento'] == t and not s['q1'] - 1.5 * s['iqr'] <= r['tempo_segundos'] <= s['q3'] + 1.5 * s['iqr']:
                outliers.append({k: r[k] for k in ('integrante', 'kata', 'tratamento', 'tempo_segundos')})
    manifest = {'fontes': {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
                          for p in [ROOT / 'lab02/data/trials.csv', ROOT / 'lab02/data/static_metrics.csv']},
                'quartis': 'interpolação linear inclusiva (tipo 7)', 'n_trials': len(rows),
                'resumo': stats, 'testes': tests, 'outliers_tempo_tukey_1_5_iqr': outliers,
                'metricas_conferidas_com_radon': 18}
    (OUT / 'resultados_relatorio.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + '\n')
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.75)
    sec.header_distance = sec.footer_distance = Inches(.25)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption'):
        style = d.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    d.styles['Caption'].font.bold = False
    normal = d.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [('Title', 23), ('Heading 1', 15), ('Heading 2', 12), ('Caption', 9)]:
        d.styles[name].font.size = Pt(size)
    d.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    d.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    d.styles['Heading 2'].paragraph_format.space_before = Pt(9)
    d.styles['Caption'].paragraph_format.space_after = Pt(7)
    header = sec.header.paragraphs[0]
    header.text = 'PUC MINAS  |  LABORATÓRIO DE EXPERIMENTAÇÃO DE SOFTWARE  |  LAB02'
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Lab02  •  Relatório final  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    d.core_properties.title = 'Assistentes de IA e codificação manual'
    d.core_properties.author = 'Gabriel Ferreira Amaral'
    d.core_properties.subject = 'Relatório final do Lab02'

    text(d, 'Assistentes de IA e\ncodificação manual', 'Title')
    text(d, 'Relatório final do experimento controlado', 'Subtitle')
    text(d, 'Grupo: Gabriel Ferreira Amaral, Gabriel Santiago e Gabriel Tavares.\n'
         'Responsável pelo relatório: Gabriel Ferreira Amaral.\n'
         'Professor: Danilo Maia  •  Engenharia de Software  •  Setembro de 2026', 'Caption')
    text(d, 'Objetivo e resultado principal', 'Heading 1')
    text(d, 'Avaliamos a associação entre o uso de ChatGPT e o tempo de resolução, o sucesso nos testes '
         'e a estrutura de seis katas Python. Os tempos registrados foram menores com IA, mas a análise '
         'pareada de três participantes não alcançou significância estatística (p = 0,125). '
         'Todos os trials passaram nos testes, e a mediana da complexidade foi igual nos dois tratamentos.')
    text(d, 'Hipóteses e métricas', 'Heading 1')
    table(d, ['Questão', 'Hipótese alternativa', 'Métrica escolhida'], [
        ['RQ1 Tempo', 'Menor tempo com IA', 'Segundos até todos os testes passarem'],
        ['RQ2 Defeitos', 'Maior sucesso com IA', 'Testes passando / total'],
        ['RQ3 Estrutura', 'Complexidade diferente com IA', 'CC média por função; LOC como controle']])
    text(d, 'As hipóteses nulas estabelecem diferenças centradas em zero entre os tratamentos. '
         'A taxa de sucesso permite comparar katas com cinco ou seis testes; LOC contextualiza '
         'o tamanho das soluções. Reportamos mediana e intervalo interquartil (IQR) pela amostra pequena.')
    text(d, 'Desenho e execução', 'Heading 1')
    text(d, 'Três participantes resolveram seis katas cada: 18 trials, nove com IA e nove sem IA. '
         'Cada pessoa realizou três de cada tratamento, em katas distintos. Os exercícios autorais '
         'abrangem intervalos, dicionários, logs, vagas livres, checksum e pontuação de frames. '
         'O limite foi 35 minutos por trial; nenhum registro atingiu a censura.')
    text(d, 'A sequência registrada foi IA → sem IA para Ferreira, sem IA → IA para Santiago e '
         'alternada para Tavares. O protocolo previa Python 3, pytest, Radon e ChatGPT web gratuito '
         'como único assistente. A versão do modelo e a IDE não foram registradas, limitando a replicação exata.')
    text(d, 'Para inferência, usamos as medianas por pessoa e tratamento: três pares independentes, '
         'sem tratar os nove trials de cada grupo como nove pares. Aplicamos Wilcoxon por permutação '
         'exaustiva de sinais, unilateral em RQ1/RQ2 e bilateral em RQ3, com alfa de 5%.')

    d.add_page_break()
    text(d, 'Tempo e sucesso nos testes', 'Heading 1')
    text(d, 'RQ1 Tempo de resolução', 'Heading 2')
    s = stats['tempo_segundos']
    text(d, f"A mediana dos nove trials foi {fmt(s['com_ia']['mediana'])} s com IA e "
         f"{fmt(s['sem_ia']['mediana'])} s sem IA; os IQRs foram {fmt(s['com_ia']['iqr'])} s e "
         f"{fmt(s['sem_ia']['iqr'])} s, respectivamente. A figura compara as medianas individuais "
         'que entram no teste pareado.')
    picture(d, 'rq1_tempo_pareado', 'Figura 1. Medianas de três trials por ponto. Cada linha liga a mesma pessoa '
            'nos dois tratamentos. Fonte: trials.csv; n = 3 participantes, 18 trials.')
    text(d, 'Os três participantes tiveram menor mediana com IA. O teste resultou em W+ = 0 e '
         'p = 0,125: não rejeitamos H0. Com apenas três diferenças não nulas, o menor p unilateral '
         'possível é 1/8 = 0,125. A separação visual não permite afirmar um efeito generalizável.')
    text(d, 'RQ2 Qualidade funcional', 'Heading 2')
    picture(d, 'rq2_sucesso', 'Figura 2. Taxa de sucesso final em escala completa de 0 a 100%. Os nove pontos '
            'de cada grupo coincidem. Fonte: trials.csv; n = 9 trials por tratamento.')
    text(d, 'Os 18 trials atingiram 100% de sucesso e zero testes falhando. As diferenças das '
         'medianas individuais foram [0, 0, 0]; Wilcoxon não tem diferenças não nulas para testar. '
         'Esse efeito teto impede avaliar redução de defeitos. Passar nos testes disponíveis '
         'não demonstra ausência de defeitos ou equivalência entre os tratamentos.')
    text(d, 'Leitura conjunta', 'Heading 2')
    text(d, 'Nesta amostra, a diferença observável está no tempo registrado. O encerramento no '
         'primeiro green torna a taxa final pouco discriminativa quando todos concluem. '
         'Uma nova coleta deve registrar também a primeira tentativa e o número de tentativas até green.')

    d.add_page_break()
    text(d, 'Estrutura das soluções', 'Heading 1')
    text(d, 'RQ3 Complexidade com controle de tamanho', 'Heading 2')
    text(d, 'Conferimos as 18 medições de complexidade e LOC com Radon 6.0.1 nos arquivos finais '
         'preservados por participante. A mediana da CC foi 4 nos dois tratamentos; a LOC física '
         'foi 23 com IA e 17 sem IA. O número de linhas, isoladamente, não indica maior complexidade.')
    picture(d, 'rq3_estrutura', 'Figura 3. Cada ponto representa um trial; deslocamento horizontal revela '
            'valores repetidos. Traço preto = mediana; faixa vertical = Q1 a Q3. '
            'Fonte: static_metrics.csv, conferido nas soluções; n = 9 por tratamento.')
    vals = []
    for metric, label in [('complexidade_media', 'CC média por função'), ('loc', 'Linhas físicas LOC')]:
        vals.append([label] + [f"{fmt(stats[metric][t]['mediana'], 0)} / {fmt(stats[metric][t]['iqr'], 0)}"
                              for t in ('com_ia', 'sem_ia')])
    table(d, ['Métrica', 'Com IA mediana / IQR', 'Sem IA mediana / IQR'], vals)
    text(d, 'As medianas individuais de CC também foram iguais: Ferreira 4/4, Santiago 5/5 e '
         'Tavares 4/4 (com/sem IA). A análise bilateral das medianas não possui diferenças não '
         'nulas para Wilcoxon. A conclusão é descritiva: não observamos mudança na mediana '
         'de CC, sem demonstrar equivalência estrutural.')
    text(d, 'LOC funciona como controle descritivo; não ajustamos um modelo causal nem dividimos '
         'CC automaticamente por LOC. A duplicação não foi medida: o grupo optou por CC e LOC '
         'em funções curtas. Portanto, os resultados não respondem à parcela de duplicação da RQ3. '
         'O índice de manutenibilidade, opcional, foi mantido fora desta síntese.')
    text(d, 'Convenção estatística', 'Heading 2')
    text(d, 'Usamos quartis por interpolação linear inclusiva (tipo 7) em todas as métricas. '
         'Essa padronização coincide com o dashboard atual; o relatório inicial de RQ3 usava '
         'outra convenção de quartis, por isso seus IQRs de LOC diferem dos apresentados aqui.')

    d.add_page_break()
    text(d, 'Discussão e conclusão', 'Heading 1')
    text(d, 'Validade e revisão dos dados', 'Heading 2')
    text(d, 'A principal restrição é o tamanho amostral: três participantes não permitem '
         'confirmar o efeito de tempo a 5% pelo teste exato adotado. Katas diferentes em cada '
         'tratamento, aprendizado, ordem e familiaridade com IA podem explicar parte das diferenças. '
         'A simetria das diferenças, pressuposto do teste, não pode ser avaliada com segurança.')
    text(d, 'Pela regra de 1,5 × IQR, os tempos sem IA de 122,79 s, 245,54 s, 664,54 s e 908,87 s '
         'são observações extremas. Todos foram mantidos: ser extremo não comprova erro, '
         'especialmente com nove trials por tratamento. Nenhum tempo com IA ultrapassou esses limites.')
    text(d, 'O README do dashboard identifica os CSVs como fonte. No commit b09955d, três '
         'tempos sem IA de Ferreira foram alterados para 426,47 s, 473,86 s e 521,25 s, '
         'sem explicação específica da remedição no README. Usamos a versão atual e preservamos '
         'essa ressalva de procedência. Os tempos com IA de poucos segundos também exigem '
         'cuidado ao interpretá-los como duração integral de programação.')
    text(d, 'Conclusão', 'Heading 2')
    text(d, 'O experimento registra tempos menores com IA, mas não sustenta uma conclusão '
         'inferencial de redução de tempo. Não distingue os tratamentos quanto a defeitos '
         'nos testes finais ou à mediana de complexidade. Para ampliar a evidência, são '
         'necessários mais participantes, registro verificável de início e fim e tarefas '
         'que permitam observar falhas antes do green.')
    text(d, 'Reprodução e rastreabilidade', 'Heading 2')
    text(d, 'Dados: lab02/data/trials.csv e static_metrics.csv. Soluções: lab02/trials/. '
         'Scripts: analyze_rq1_rq2.py, static_metrics.py e scripts/gerar_relatorio_lab02.py. '
         'O arquivo resultados_relatorio.json registra os hashes das fontes, quartis, testes '
         'e observações extremas. As dependências e os comandos estão em relatorio/lab02/README.md.')
    text(d, 'O ambiente de coleta foi descrito pelo grupo como Python 3.14; a geração deste '
         'relatório usa Python 3.12, SciPy 1.18.1 e Radon 6.0.1. Os timestamps e durações '
         'originais não foram reconstituídos. A interpretação do Wilcoxon segue a documentação '
         'do SciPy, e o desenho segue o enunciado do Lab02.')
    link(d, 'Repositório', 'https://github.com/Druitti/Laboratorio-de-Desenvolvimento-de-Software')
    link(d, 'GitHub Projects', 'https://github.com/users/Druitti/projects/5')
    link(d, 'Entrega do relatório', 'https://github.com/Druitti/Laboratorio-de-Desenvolvimento-de-Software/issues/40')
    link(d, 'Referência estatística', 'https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.wilcoxon.html')
    d.save(OUT / 'Lab02_Relatorio_Final.docx')
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    FIG.mkdir(parents=True, exist_ok=True)
    data = load_data()
    figures(data)
    build(data)
